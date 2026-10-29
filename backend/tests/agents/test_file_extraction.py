"""
Run:
uv run pytest tests/agents/test_file_extraction.py -s
"""

import io

import docx
import pytest

from services.file_extraction import extract_text
from tests.conftest import SAMPLE_PDF


def test_extract_text_from_pdf(sample_pdf_bytes):
    text = extract_text(sample_pdf_bytes, SAMPLE_PDF.name)

    print(f"\n--- Extracted text from {SAMPLE_PDF.name} ({len(text)} chars) ---")
    print(text)

    assert isinstance(text, str)
    assert text.strip()



def test_extract_text_from_docx():
    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("   ")
    doc.add_paragraph("Python developer")
    buf = io.BytesIO()
    doc.save(buf)

    assert extract_text(buf.getvalue(), "resume.docx") == "Jane Doe\n\nPython developer"


def test_unsupported_file_type_raises():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(b"hello", "resume.txt")


def test_empty_document_raises():
    buf = io.BytesIO()
    docx.Document().save(buf)

    with pytest.raises(ValueError, match="no extractable text"):
        extract_text(buf.getvalue(), "empty.docx")
